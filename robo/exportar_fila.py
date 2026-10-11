"""
Gera a planilha com a fila de propostas que aguardam a decisão do gerente.
Raia "Gerente de Negócios" do TO-BE: a tarefa de usuário "Analisar e decidir proposta".

Uso: python robo/exportar_fila.py
"""
import logging
import sqlite3
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from config_log import configurar_log
from ler_propostas import CAMINHO_BANCO

PASTA_FILA = Path(__file__).parent.parent / "fila"
ARQUIVO_FILA = PASTA_FILA / "fila_gerente.xlsx"

COLUNAS = ["Proposta", "Associado", "Agência", "Valor (R$)", "Prazo (meses)",
           "Comprometimento", "Decisão", "Parecer"]

SQL_FILA = """
    SELECT p.id, a.nome, a.agencia, p.valor, p.prazo_meses,
           (SELECT l.detalhe FROM log_etapa l
            WHERE l.proposta_id = p.id AND l.etapa = 'CALCULO_COMPROMETIMENTO'
            ORDER BY l.id DESC LIMIT 1) AS comprometimento
    FROM proposta p
    JOIN associado a ON a.id = p.associado_id
    WHERE p.etapa_atual = 'GERENTE'
    ORDER BY p.data_solicitacao
"""

logger = logging.getLogger("fila")


def montar_planilha(propostas):
    """Cria a planilha com as propostas e as colunas de decisão em branco."""
    planilha = Workbook()
    aba = planilha.active
    aba.title = "Fila do gerente"
    aba.append(COLUNAS)

    for proposta in propostas:
        aba.append([
            proposta["id"], proposta["nome"], proposta["agencia"], proposta["valor"],
            proposta["prazo_meses"], proposta["comprometimento"] or "-", None, None,
        ])

    # Cabeçalho em negrito e colunas a preencher destacadas
    for celula in aba[1]:
        celula.font = Font(bold=True)
    destaque = PatternFill("solid", fgColor="FFF2CC")
    for linha in range(2, len(propostas) + 2):
        aba.cell(row=linha, column=7).fill = destaque
        aba.cell(row=linha, column=8).fill = destaque

    # Lista suspensa: o gerente só consegue escolher valores válidos
    validacao = DataValidation(
        type="list",
        formula1='"APROVADA,REPROVADA"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Decisão inválida",
        error="Escolha APROVADA ou REPROVADA na lista.",
    )
    aba.add_data_validation(validacao)
    validacao.add(f"G2:G{len(propostas) + 1}")

    for coluna, largura in zip("ABCDEFGH", [10, 22, 18, 12, 14, 16, 14, 50]):
        aba.column_dimensions[coluna].width = largura

    return planilha


def main():
    configurar_log()

    conexao = sqlite3.connect(CAMINHO_BANCO)
    conexao.row_factory = sqlite3.Row
    try:
        propostas = conexao.execute(SQL_FILA).fetchall()
    finally:
        conexao.close()

    PASTA_FILA.mkdir(exist_ok=True)
    try:
        montar_planilha(propostas).save(ARQUIVO_FILA)
    except PermissionError:
        logger.error("A planilha está aberta no Excel. Feche-a e rode de novo: %s", ARQUIVO_FILA)
        return 1

    logger.info("Fila do gerente gerada com %d proposta(s): %s", len(propostas), ARQUIVO_FILA)
    print(f"{len(propostas)} proposta(s) na fila do gerente.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
